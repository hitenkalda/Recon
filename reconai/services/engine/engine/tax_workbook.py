"""26AS/AIS Tax Reconciliation workbook export.

Generates a multi-sheet XLSX mirroring the Phase-8 prototype's exportXLSX:
Summary / Matched / Amount Differences / Income Mismatch / Section Mismatch /
Only in 26AS / Only in Books / Duplicates.

Row convention: each item has recordA/recordB/matchStatus/score/reasons/tax
where tax = {side, name, pan, tan, section, quarter, fy, gross, tds,
             tdsA, tdsB, grossA, grossB}.
"""
from __future__ import annotations

import io
from typing import Any, Dict, List, Optional


_INR = '#,##,##0.00'
_INR0 = '#,##,##0'
_PCT = '0.0%'

_CLEAN = {"gstin_pan", "amount", "date_tolerance", "fuzzy_invoice", "vendor_similarity"}


def _safe(v: Any) -> str:
    if v is None:
        return ""
    return str(v)


def _tax_row(it: Dict[str, Any]) -> List[Any]:
    tax = it.get("tax") or {}
    side = tax.get("side", "")
    return [
        tax.get("pan") or "",
        tax.get("tan") or "",
        tax.get("name") or "",
        tax.get("section") or "",
        tax.get("quarter") or "",
        tax.get("fy") or "",
        tax.get("gross") or 0.0,
        tax.get("tds") or 0.0,
        it.get("matchStatus") or "",
        round(it.get("score") or 0.0, 4),
        "; ".join(it.get("reasons") or []),
        side,
    ]


_TAX_HEADERS = ["PAN", "TAN", "Name", "Section", "Quarter", "FY",
                "Gross Amount", "TDS Amount", "Status", "Score", "Reasons", "Side"]


def build_workbook(analysis: Dict[str, Any], client_name: str = "Client") -> Any:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    header_fill = PatternFill("solid", fgColor="1A2B4A")
    accent_fill = PatternFill("solid", fgColor="2F6CC8")
    warn_fill = PatternFill("solid", fgColor="F4CCCC")
    header_font = Font(color="FFFFFF", bold=True, size=10, name="Calibri")
    body_font = Font(size=10, name="Calibri")
    bold_font = Font(size=10, bold=True, name="Calibri")
    title_font = Font(size=14, bold=True, name="Calibri", color="1A2B4A")
    subtitle_font = Font(size=10, name="Calibri", color="666666")
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    section_fill = PatternFill("solid", fgColor="E8F0FE")
    alt_fill = PatternFill("solid", fgColor="F8F9FA")

    def style_header(ws, row=1, ncols=None):
        for cell in ws[row]:
            cell.fill = header_fill
            cell.font = header_font
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def autosize(ws, widths_by_col):
        for i, w in enumerate(widths_by_col, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"

    def write_rows(ws, rows: List[List[Any]], num_cols: int, money_cols: Optional[List[int]] = None):
        money_cols = money_cols or []
        for r_idx, row in enumerate(rows):
            ws.append(row)
            last = ws.max_row
            for c in range(1, num_cols + 1):
                cell = ws.cell(row=last, column=c)
                cell.font = body_font
                cell.border = border
                if c in money_cols:
                    cell.number_format = _INR
                if r_idx % 2 == 1:
                    cell.fill = alt_fill

    items = analysis.get("items") or []
    metrics = analysis.get("metrics") or {}
    tax_m = metrics.get("tax") or {}

    wb = Workbook()

    # ── Sheet 1: Summary ──────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Summary"
    ws.merge_cells("A1:F1")
    ws["A1"] = f"26AS/AIS Tax Reconciliation — {client_name}"
    ws["A1"].font = title_font
    ws["A2"] = f"Generated: {analysis.get('generatedAt', '')}"
    ws["A2"].font = subtitle_font

    # Overview metrics
    ws["A4"] = "Overview"
    ws["A4"].font = bold_font
    overview = [
        ["Metric", "Value"],
        ["Total 26AS records", tax_m.get("totalA", metrics.get("totalA", 0))],
        ["Total Books records", tax_m.get("totalB", metrics.get("totalB", 0))],
        ["Matched pairs", tax_m.get("pairs", metrics.get("matched", 0))],
        ["Match rate", f"{tax_m.get('matchRate', 0)}%"],
        ["TDS per 26AS", tax_m.get("tds26", 0)],
        ["TDS per Books", tax_m.get("tdsBooks", 0)],
        ["Net TDS difference", tax_m.get("netDiff", 0)],
        ["TDS missing in books (credit at risk)", tax_m.get("missingInBooks", {}).get("tds", 0)],
        ["TDS missing in 26AS (unfiled deductor)", tax_m.get("missingIn26AS", {}).get("tds", 0)],
        ["Amount differences", tax_m.get("amountDiff", {}).get("count", 0)],
        ["Income mismatches", tax_m.get("incomeMismatch", {}).get("count", 0)],
        ["Section mismatches", tax_m.get("sectionMismatch", {}).get("count", 0)],
        ["Duplicates", tax_m.get("duplicates", {}).get("count", 0)],
        ["High risk", tax_m.get("highRisk", {}).get("count", 0)],
    ]
    for row in overview:
        ws.append(row)
        last = ws.max_row
        for c in range(1, 3):
            cell = ws.cell(row=last, column=c)
            cell.font = bold_font if last == 2 else body_font
            cell.border = border
            if last > 2 and c == 2 and isinstance(cell.value, (int, float)):
                cell.number_format = _INR
    style_header(ws, row=2)
    autosize(ws, [42, 22])

    # Section summary
    sec_sum = metrics.get("sectionSummary") or []
    if sec_sum:
        ws.append([])
        row_start = ws.max_row + 1
        ws.cell(row=row_start, column=1, value="Section Summary").font = bold_font
        ws.append(["Section", "Count", "Gross", "TDS", "Matched", "Missing Books",
                    "Missing 26AS", "Amt Diff", "Income Mismatch", "Section Mismatch", "Duplicates"])
        style_header(ws, row=ws.max_row)
        for s in sec_sum:
            ws.append([s.get("section", "-"), s.get("count", 0), s.get("gross", 0), s.get("tds", 0),
                        s.get("matched", 0), s.get("missingInBooks", 0), s.get("missingIn26AS", 0),
                        s.get("amountDiff", 0), s.get("incomeMismatch", 0), s.get("sectionMismatch", 0),
                        s.get("duplicates", 0)])
            last = ws.max_row
            for c in range(1, 12):
                cell = ws.cell(row=last, column=c)
                cell.font = body_font
                cell.border = border
                if c in (3, 4):
                    cell.number_format = _INR
        autosize(ws, [32, 10, 14, 14, 10, 14, 14, 10, 16, 16, 12])

    # Deductor summary
    ded_sum = metrics.get("deductorSummary") or []
    if ded_sum:
        ws.append([])
        ws.cell(row=ws.max_row + 1, column=1, value="Deductor Summary (Top 20)").font = bold_font
        ws.append(["Deductor", "PAN", "TAN", "Count", "Gross", "TDS", "Matched", "Missing"])
        style_header(ws, row=ws.max_row)
        for d in ded_sum[:20]:
            ws.append([d.get("deductor", ""), d.get("pan") or "", d.get("tan") or "",
                        d.get("count", 0), d.get("gross", 0), d.get("tds", 0),
                        d.get("matched", 0), d.get("missing", 0)])
            last = ws.max_row
            for c in range(1, 9):
                cell = ws.cell(row=last, column=c)
                cell.font = body_font
                cell.border = border
                if c in (5, 6):
                    cell.number_format = _INR
        autosize(ws, [30, 12, 14, 10, 14, 14, 10, 10])

    # ── Sheet 2: Matched ──────────────────────────────────────────────────────
    ws2 = wb.create_sheet("Matched")
    ws2.append(_TAX_HEADERS)
    style_header(ws2)
    matched = [it for it in items if it.get("matchStatus") in _CLEAN]
    write_rows(ws2, [_tax_row(it) for it in matched], len(_TAX_HEADERS), money_cols=[7, 8])
    autosize(ws2, [12, 14, 30, 10, 8, 8, 14, 14, 16, 8, 40, 10])

    # ── Sheet 3: Amount Differences ───────────────────────────────────────────
    ws3 = wb.create_sheet("Amount Differences")
    ws3.append(_TAX_HEADERS + ["TDS 26AS", "TDS Books", "Difference"])
    style_header(ws3)
    amt_diff = [it for it in items if it.get("matchStatus") == "amount_tolerance"]
    for it in amt_diff:
        tax = it.get("tax") or {}
        row = _tax_row(it) + [tax.get("tdsA") or 0, tax.get("tdsB") or 0,
                               abs((tax.get("tdsA") or 0) - (tax.get("tdsB") or 0))]
        write_rows(ws3, [row], len(_TAX_HEADERS) + 3, money_cols=[7, 8, 13, 14, 15])
    if not amt_diff:
        ws3.append(["(No amount differences found)"])
    autosize(ws3, [12, 14, 30, 10, 8, 8, 14, 14, 16, 8, 40, 10, 14, 14, 14])

    # ── Sheet 4: Income Mismatch ──────────────────────────────────────────────
    ws4 = wb.create_sheet("Income Mismatch")
    ws4.append(_TAX_HEADERS + ["Gross 26AS", "Gross Books", "Gross Diff"])
    style_header(ws4)
    inc_mm = [it for it in items if it.get("matchStatus") == "income_mismatch"]
    for it in inc_mm:
        tax = it.get("tax") or {}
        ga = tax.get("grossA") or 0
        gb = tax.get("grossB") or 0
        row = _tax_row(it) + [ga, gb, abs(ga - gb)]
        write_rows(ws4, [row], len(_TAX_HEADERS) + 3, money_cols=[7, 8, 13, 14, 15])
    if not inc_mm:
        ws4.append(["(No income mismatches found)"])
    autosize(ws4, [12, 14, 30, 10, 8, 8, 14, 14, 16, 8, 40, 10, 14, 14, 14])

    # ── Sheet 5: Section Mismatch ─────────────────────────────────────────────
    ws5 = wb.create_sheet("Section Mismatch")
    ws5.append(_TAX_HEADERS + ["Section 26AS", "Section Books"])
    style_header(ws5)
    sec_mm = [it for it in items if it.get("matchStatus") == "section_mismatch"]
    for it in sec_mm:
        tax = it.get("tax") or {}
        row = _tax_row(it) + [tax.get("sectionA") or tax.get("section") or "",
                               tax.get("sectionB") or tax.get("section") or ""]
        write_rows(ws5, [row], len(_TAX_HEADERS) + 2, money_cols=[7, 8])
    if not sec_mm:
        ws5.append(["(No section mismatches found)"])
    autosize(ws5, [12, 14, 30, 10, 8, 8, 14, 14, 16, 8, 40, 10, 14, 14])

    # ── Sheet 6: Only in 26AS (missing in books) ─────────────────────────────
    ws6 = wb.create_sheet("Only in 26AS")
    ws6.append(_TAX_HEADERS)
    style_header(ws6)
    only_26 = [it for it in items if it.get("matchStatus") == "unmatched"
               and it.get("recordA") and not it.get("recordB")]
    write_rows(ws6, [_tax_row(it) for it in only_26], len(_TAX_HEADERS), money_cols=[7, 8])
    if not only_26:
        ws6.append(["(All 26AS entries matched with books)"])
    autosize(ws6, [12, 14, 30, 10, 8, 8, 14, 14, 16, 8, 40, 10])

    # ── Sheet 7: Only in Books (missing in 26AS) ─────────────────────────────
    ws7 = wb.create_sheet("Only in Books")
    ws7.append(_TAX_HEADERS)
    style_header(ws7)
    only_b = [it for it in items if it.get("matchStatus") == "unmatched"
              and not it.get("recordA") and it.get("recordB")]
    write_rows(ws7, [_tax_row(it) for it in only_b], len(_TAX_HEADERS), money_cols=[7, 8])
    if not only_b:
        ws7.append(["(All booked TDS found in 26AS)"])
    autosize(ws7, [12, 14, 30, 10, 8, 8, 14, 14, 16, 8, 40, 10])

    # ── Sheet 8: Duplicates ───────────────────────────────────────────────────
    ws8 = wb.create_sheet("Duplicates")
    ws8.append(_TAX_HEADERS)
    style_header(ws8)
    dups = [it for it in items if it.get("matchStatus") == "duplicate"]
    write_rows(ws8, [_tax_row(it) for it in dups], len(_TAX_HEADERS), money_cols=[7, 8])
    if not dups:
        ws8.append(["(No duplicates found)"])
    autosize(ws8, [12, 14, 30, 10, 8, 8, 14, 14, 16, 8, 40, 10])

    return wb


def workbook_bytes(analysis: Dict[str, Any], client_name: str = "Client") -> bytes:
    buf = io.BytesIO()
    build_workbook(analysis, client_name).save(buf)
    return buf.getvalue()
